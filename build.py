"""Build the training dashboard site from the HSI Training Summary file.

1. Gets the Training Summary Excel file from SharePoint through the
   Power Automate flow at TRAINING_SUMMARY_URL (the flow answers with the
   file itself). For a local test, pass a file path instead:
       python build.py "Training Summary.xlsx"
2. Turns the rows into the dashboard's data format.
3. If the DASHBOARD_PASSCODE secret is set, locks the data with it (AES-GCM)
   so the page asks for the passcode. Without it, the data is published open.
4. Writes the finished site to site/.
"""
import io
import re
import os
import sys
import json
import base64
import shutil
from datetime import datetime, date, timezone

import requests
from openpyxl import load_workbook
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "site")
ITERATIONS = 200_000
NO_EXPIRATION_YEAR = 2050  # expire dates this far out mean "doesn't expire"

# Dashboard column -> header names the export might use (lowercase, no spaces/punctuation).
COLUMNS = {
    "Employee ID": ["employeeid", "empid", "employeenumber", "eeid", "id"],
    "Name": ["name", "employeename", "fullname", "membername"],
    "Role": ["role", "jobtitle", "title", "position"],
    "Location": ["location", "team", "teams", "site"],
    "Course Name": ["coursename", "course", "training", "trainingname", "coursetitle"],
    "Score": ["score"],
    "Date Assigned": ["dateassigned", "assigned", "assigneddate", "assignmentdate"],
    "Date Completed": ["datecompleted", "completed", "completeddate", "completiondate"],
    "Date Expire": ["dateexpire", "expire", "expires", "expiredate", "expirationdate", "dateexpires"],
    "Status": ["status", "coursestatus", "completionstatus"],
}
REQUIRED = ["Name", "Course Name", "Date Expire", "Status"]


# HSI system and test logins that aren't real employees (DoneSafe, Test Student, ...).
SYSTEM_ACCOUNT = re.compile(r"donesafe|^\s*test\b|\btest\s+student\b|\bstudent\s*,?\s*test\b|limited access", re.I)


def norm(text):
    return "".join(ch for ch in str(text or "").lower() if ch.isalnum())


def get_file():
    if len(sys.argv) > 1:
        with open(sys.argv[1], "rb") as f:
            return f.read()
    secrets = json.loads(os.environ.get("ALL_SECRETS") or "{}")
    url = secrets.get("TRAINING_SUMMARY_URL") or os.environ.get("TRAINING_SUMMARY_URL")
    if not url:
        raise SystemExit("Missing GitHub secret: TRAINING_SUMMARY_URL")
    print("Getting the Training Summary file from SharePoint...")
    res = requests.post(url, json={}, timeout=180)
    if res.status_code != 200:
        raise SystemExit(f"The SharePoint flow returned {res.status_code}: {res.text[:300]}")
    body = res.content
    if not body.startswith(b"PK"):
        # The flow may have answered with JSON holding the file as base64.
        try:
            j = res.json()
            b = j.get("$content") or j.get("fileContent") or j.get("body", {}).get("$content")
            body = base64.b64decode(b)
        except Exception:
            raise SystemExit("The SharePoint flow didn't send back an Excel file. "
                             f"First bytes: {res.content[:120]!r}")
    print(f"Got {len(body):,} bytes.")
    return body


def iso(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        d = value.date()
    elif isinstance(value, date):
        d = value
    else:
        text = str(value).strip()
        d = None
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y %I:%M %p"):
            try:
                d = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                continue
        if d is None:
            return None
    return d.isoformat()


def read_rows(xlsx_bytes):
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    best = None
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        for i, row in enumerate(rows):
            if i > 15:
                break
            headers = [norm(c) for c in row]
            found = {}
            for col, names in COLUMNS.items():
                # Exact column name first, then the other names in order of preference.
                for want in [norm(col)] + names:
                    if want in headers:
                        found[col] = headers.index(want)
                        break
            if all(c in found for c in REQUIRED):
                best = (ws, i, found, [str(c or "") for c in row])
                break
        if best:
            break
    if not best:
        sheet_info = []
        for ws in wb.worksheets:
            first = next(ws.iter_rows(values_only=True, max_row=1), ())
            sheet_info.append(f"{ws.title}: {[c for c in first if c]}")
        raise SystemExit("Couldn't find the header row (need Name, Course Name, Date Expire, Status). "
                         "Sheets and first rows: " + " | ".join(sheet_info))

    ws, header_row, found, header_text = best
    print(f"Using sheet '{ws.title}', headers on row {header_row + 1}: "
          + ", ".join(f"{c} <- '{header_text[i]}'" for c, i in found.items()))
    cols = list(COLUMNS)
    out = []
    for row in ws.iter_rows(min_row=header_row + 2, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue
        rec = []
        for col in cols:
            idx = found.get(col)
            v = row[idx] if idx is not None and idx < len(row) else None
            if col.startswith("Date"):
                v = iso(v)
                if col == "Date Expire" and v and int(v[:4]) >= NO_EXPIRATION_YEAR:
                    v = None
            elif isinstance(v, str):
                v = v.strip() or None
            elif isinstance(v, float) and v.is_integer():
                v = int(v)
            rec.append(v)
        if not rec[cols.index("Name")]:
            continue
        out.append(rec)
    return cols, out


def encrypt(obj, passcode):
    salt, iv = os.urandom(16), os.urandom(12)
    key = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=ITERATIONS).derive(passcode.encode())
    data = AESGCM(key).encrypt(iv, json.dumps(obj, separators=(",", ":")).encode(), None)
    b = lambda x: base64.b64encode(x).decode()
    return {"v": 1, "iter": ITERATIONS, "salt": b(salt), "iv": b(iv), "data": b(data)}


TAPPI_COLUMNS = {"first": ["firstname"], "last": ["lastname"], "id": ["studentid"], "course": ["description", "course", "coursename"],
                 "date": ["classdate", "datecompleted"], "exp": ["expirationdate", "expiredate"], "grade": ["grade", "result"]}


def fetch_xlsx(secret_name, label):
    """Post to a Power Automate flow that answers with an Excel file. None if the secret isn't set."""
    secrets = json.loads(os.environ.get("ALL_SECRETS") or "{}")
    url = secrets.get(secret_name) or os.environ.get(secret_name)
    if not url:
        return None
    res = requests.post(url, json={}, timeout=180)
    body = res.content
    if res.status_code == 200 and not body.startswith(b"PK"):
        try:
            j = res.json()
            body = base64.b64decode(j.get("$content") or j.get("fileContent") or j.get("body", {}).get("$content"))
        except Exception:
            pass
    if res.status_code != 200 or not body.startswith(b"PK"):
        print(f"::warning title={label}::The {label} flow didn't send back an Excel file (status {res.status_code}); skipping it.")
        return None
    return body


def read_tappi(xlsx_bytes):
    """TAPPISAFE Student Attendance export -> compact rows for the Plant Sites tab."""
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    for h, row in enumerate(rows[:15]):
        headers = [norm(c) for c in row]
        found = {k: next((headers.index(n) for n in names if n in headers), None) for k, names in TAPPI_COLUMNS.items()}
        if found["course"] is not None and found["last"] is not None:
            break
    else:
        print("::warning title=TAPPI::Couldn't find the header row in the TAPPI file; skipping the Plant Sites tab.")
        return None
    cols = list(TAPPI_COLUMNS)
    out = []
    for r in rows[h + 1:]:
        g = lambda k: r[found[k]] if found[k] is not None and found[k] < len(r) else None
        if not g("course") or not g("last"):
            continue
        out.append([str(g("first") or "").strip(), str(g("last") or "").strip(), str(g("id") or "").strip(),
                    str(g("course")).strip(), iso(g("date")), iso(g("exp")), str(g("grade") or "").strip()])
    print(f"TAPPI: {len(out):,} course records.")
    return {"cols": cols, "rows": out}


def read_refresher(xlsx_bytes):
    """HSI Training Requirement Status > By Team/Member - All -> compact rows for the Refresher Training view."""
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    h = next((i for i, r in enumerate(rows[:15]) if r and "Requirement" in [str(v or "").strip() for v in r]
              and "First Name" in [str(v or "").strip() for v in r]), None)
    if h is None:
        print("::warning title=Refresher::Couldn't find the header row in the refresher file; skipping it.")
        return None
    hdr = [str(v or "").strip() for v in rows[h]]
    col = lambda n: hdr.index(n) if n in hdr else None
    c = {k: col(n) for k, n in (("first", "First Name"), ("last", "Last Name"), ("team", "Teams"), ("req", "Requirement"),
                                 ("status", "Compliance Status"), ("last_done", "Last Completion Date"),
                                 ("window", "Start Date for Next Due"), ("due", "Due Date"))}
    reqs, out, skipped = [], [], 0
    for r in rows[h + 1:]:
        g = lambda k: r[c[k]] if c[k] is not None and c[k] < len(r) else None
        if not g("req") or not g("last"):
            continue
        last, emp_id = str(g("last")).strip(), ""
        if "," in last:  # HSI puts the login ID after the last name: "Ferguson,3LDUWVKMI"
            last, emp_id = [x.strip() for x in last.split(",", 1)]
        req = str(g("req")).strip()
        if SYSTEM_ACCOUNT.search(f"{g('first') or ''} {last}") or SYSTEM_ACCOUNT.search(str(g("last"))):
            skipped += 1
            continue
        if req not in reqs:
            reqs.append(req)
        out.append([emp_id, (str(g("first") or "").strip() + " " + last).strip(), str(g("team") or "").strip(),
                    reqs.index(req), str(g("status") or "").strip(), iso(g("last_done")), iso(g("window")), iso(g("due"))])
    med = re.compile(r"respirator fit|fit[- ]?test|heavy metal|blood panel|medical exam|physical exam|\bPFTs?\b|pulmonary function|audiogram|hearing test|rapid meq|respirator clearance", re.I)
    print("::notice title=Medical requirements::" + "; ".join(q for q in reqs if med.search(q)))
    lvl = re.compile(r"\(\s*(1\s*y(r|ear)?|4\s*m|8\s*m)\s*\)|tech\s*ii\s*to\s*tech\s*iii\s*skill\s*assessment", re.I)
    print("::notice title=Level Up requirements::" + "; ".join(q for q in reqs if lvl.search(q) and not med.search(q)))
    print("::notice title=All refresher requirements::" + "; ".join(reqs)[:3500])
    print(f"::notice title=Refresher::{len(out):,} requirement rows, {len(reqs)} requirements; left out {skipped} rows for test accounts.")
    return {"reqs": reqs, "rows": out}


# SharePoint fields that come with every list item but aren't part of the list itself.
SP_SKIP = re.compile(r"^(@odata|odata|{|ItemInternalId|ID$|Id$|Modified|Created|Author|Editor|_|OData__|ComplianceAsset|GUID|Attachments|ContentType|"
                     r"Identifier|IsFolder|Thumbnail|LinkTitle|FilenameWithExtension|DocIcon|AppAuthor|AppEditor|ItemChildCount|FolderChildCount|"
                     r"FileSystemObjectType|ServerRedirected|EditorId|AuthorId|AttachmentFiles)", re.I)


# Friendlier names for the list's shortened column names.
NICE_COL = {"Plant/ Mill Si": "Site", "Number of Site_x0020": "Number of site trainings", "Site Specific Traini": "Site specific training",
            "Basic Training Requi": "Basic training", "StudentLink": "Training link", "RandomSiteRequirements": "Other site requirements",
            "FIT Test Requirement": "Fit test", "Background Requirements": "Background check", "Drug Screen Type": "Drug screen"}
SECRET_COL = re.compile(r"password|admin|login|username|credential", re.I)
DIRECTIONS_COL = re.compile(r"directions", re.I)
# Any sentence or line of the directions that talks about signing in is left out (the site is public).
CRED_TEXT = re.compile(r"pass\s*word|passwd|\bpwd\b|\bpw\b|user\s*name|\buser\s*id\b|log\s*-?\s*in|sign\s*-?\s*in|credential|access\s*code|\bpin\b|"
                       r"\bcode\s*(is|:)|@[A-Za-z]*\d|\b[A-Za-z]+\d{2,}\b.*\bexample\b", re.I)


def safe_directions(text):
    """Keep the directions, minus any sentence that mentions passwords, usernames, logins or codes."""
    parts = re.split(r"(?<=[.!?])\s+|\n+", str(text))
    kept = [p.strip() for p in parts if p.strip() and not CRED_TEXT.search(p)]
    return " ".join(kept), len([p for p in parts if p.strip()]) - len(kept)


def previous_site_requirements():
    """The site requirements currently on the live dashboard, so a slow or failed flow run doesn't blank the tab."""
    try:
        r = requests.get("https://envirovac-vecta.github.io/Training-dashboard/data.json", timeout=60)
        prev = r.json().get("siteReqs") if r.status_code == 200 else None
        if prev and prev.get("rows"):
            print(f"::notice title=Site requirements::Kept {len(prev['rows'])} sites from the last good refresh.")
            return prev
    except Exception:
        pass
    return None


# ---------- TAPPISAFE course lengths (course_lengths.json, from TappiSafe_Course_Lengths.xlsx) ----------
def _ln(text):
    text = re.sub(r"\([^)]*\)", " ", str(text or "").lower())
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def load_lengths():
    try:
        with open(os.path.join(HERE, "course_lengths.json")) as f:
            return json.load(f)
    except Exception as e:
        print(f"::warning title=Course lengths::course_lengths.json not read ({e}); no course times shown.")
        return []


LENGTHS = load_lengths()
_BY_CODE = {c["code"].upper(): c["min"] for c in LENGTHS if c.get("code")}
_BY_NAME = {_ln(c["course"]): c["min"] for c in LENGTHS}
HAS_TIME = re.compile(r"\(\s*\d[^)]*\b(min|mins|minutes|hr|hrs|hour|hours)\b[^)]*\)\s*$", re.I)


def course_minutes(name):
    """Minutes for a TAPPISAFE course name (same matching as courseMin() in index.html), or None."""
    for code in re.findall(r"\(([A-Za-z0-9]{5,})\)", str(name or "")):
        if code.upper() in _BY_CODE:
            return _BY_CODE[code.upper()]
    n = _ln(name)
    if not n:
        return None
    for cand in (n, re.sub(r"^tappi ?safe ", "", n)):
        if cand in _BY_NAME:
            return _BY_NAME[cand]
    best = None
    for cn, m in _BY_NAME.items():
        # a cut-off name ("...Mill and PSM Orientatio") or a line that names the course inside other words
        if (len(n) >= 15 and cn.startswith(n)) or (len(cn) >= 15 and f" {cn} " in f" {n} "):
            if best is None or len(cn) > best[0]:
                best = (len(cn), m)
    return best[1] if best else None


def fmt_minutes(m):
    h, mm = divmod(int(m), 60)
    return " ".join(x for x in (f"{h} hr" if h else "", f"{mm} min" if mm else "") if x)


def list_time_report(rows):
    """Print, per site, the Site specific training text with each course's time added in (), ready to paste into the list."""
    out, missing = [], set()
    for r in rows:
        lines = [x.strip() for x in re.split(r"\r?\n|;\s*(?=\S)", str(r.get("Site specific training") or "")) if x.strip()]
        if not lines:
            continue
        new = []
        for x in lines:
            m = None if HAS_TIME.search(x) else course_minutes(x)
            if m is None and not HAS_TIME.search(x):
                missing.add(x)
            new.append(f"{x} ({fmt_minutes(m)})" if m is not None else x)
        if new != lines:
            out.append({"site": r.get("Site", ""), "old": "\n".join(lines), "new": "\n".join(new)})
    print("::notice title=List times::" + json.dumps({"changes": out, "noMatch": sorted(missing)})[:60000])


def get_site_requirements():
    """Rows of the "Vecta and Envirovac Training Site Specific Requirements" Microsoft List, through the SITE_REQS_URL flow."""
    secrets = json.loads(os.environ.get("ALL_SECRETS") or "{}")
    url = secrets.get("SITE_REQS_URL") or os.environ.get("SITE_REQS_URL")
    if not url:
        return None
    # The flow walks every list item for attachments, which can be slow: wait longer and try up to 3 times.
    body, why = None, ""
    for attempt in range(3):
        try:
            res = requests.post(url, json={}, timeout=240)
            b = res.json()
            if isinstance(b, dict) and isinstance(b.get("d"), dict):   # SharePoint REST verbose format
                b = {"value": b["d"].get("results") or []}
            its = b if isinstance(b, list) else (b.get("value") or b.get("body", {}).get("value") or [])
            if res.status_code == 200 and its:
                body = b
                break
            why = f"status {res.status_code}, {len(its)} rows"
        except Exception as e:
            why = str(e)[:200]
    if body is None:
        print(f"::warning title=Site requirements::The list flow didn't send the list ({why}); keeping yesterday's site requirements.")
        return previous_site_requirements()
    items = body if isinstance(body, list) else (body.get("value") or body.get("body", {}).get("value") or [])
    # Attachments (optional): the flow can add {"attachments": [{"ID": 3, "files": [{DisplayName, AbsoluteUri}, ...]}]}.
    files_by_id = {}
    for a in (body.get("attachments") or []) if isinstance(body, dict) else []:
        files = a.get("files") or a.get("body") or []
        if isinstance(files, dict):
            files = files.get("value") or files.get("body") or []
        files_by_id[str(a.get("ID") or a.get("Id") or "")] = [
            {"name": f.get("DisplayName") or f.get("FileName") or "Attachment", "url": f.get("AbsoluteUri") or f.get("Url") or ""}
            for f in files if isinstance(f, dict) and (f.get("AbsoluteUri") or f.get("Url"))]
    dropped = 0
    hasc_filled = 0
    def clean(v):
        if isinstance(v, dict):   # choice / person / lookup columns
            return v.get("Url") or v.get("Value") or v.get("DisplayName") or v.get("Title") or v.get("Description") or ""
        if isinstance(v, list):
            return ", ".join(filter(None, (clean(x) for x in v)))
        return "" if v is None else v
    rows = []
    for it in items:
        row = {}
        item_id = str(it.get("ID") or it.get("Id") or "")
        # "Send an HTTP request to SharePoint" with $expand=AttachmentFiles puts the files on the item itself.
        af = it.get("AttachmentFiles")
        if isinstance(af, dict):
            af = af.get("results") or []
        if isinstance(af, list) and af:
            files_by_id[item_id] = [{"name": f.get("FileName") or "Attachment",
                                     "url": "https://envirovac-my.sharepoint.com" + requests.utils.quote(f["ServerRelativeUrl"]) if str(f.get("ServerRelativeUrl", "")).startswith("/") else f.get("ServerRelativeUrl", "")}
                                    for f in af if isinstance(f, dict) and f.get("ServerRelativeUrl")]
        for k, v in it.items():
            if SP_SKIP.search(k) or k.endswith("#Id") or k.endswith("@odata.type") or "#" in k:
                continue
            if SECRET_COL.search(k):   # passwords, admin logins and login directions never go on the public site
                continue
            v = clean(v)
            if DIRECTIONS_COL.search(k) and v:
                v, n = safe_directions(v)
                dropped += n
                k = "Specific directions"
            if v not in ("", None):
                row[k.replace("_x0020_", " ").replace("_x002f_", "/").replace("_x0028_", "(").replace("_x0029_", ")")] = v
        row = {NICE_COL.get(k, k): v for k, v in row.items() if k not in ("Title",)}
        if isinstance(row.get("Number of site trainings"), float) and row["Number of site trainings"].is_integer():
            row["Number of site trainings"] = int(row["Number of site trainings"])
        # HASC sites (online or in person) all train through hasc.com: fill in the link when the list leaves it blank.
        if not row.get("Training link") and re.search(r"\bHASC\b", f"{row.get('Training Platform', '')} {row.get('Training Location', '')}", re.I):
            row["Training link"] = "https://www.hasc.com/"
            hasc_filled += 1
        if files_by_id.get(item_id):
            row["Attachments"] = files_by_id[item_id]
        if row:
            rows.append(row)
    cols = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    print(f"::notice title=Site requirements list::{len(rows)} rows; columns: {', '.join(cols)[:1500]}; "
          f"{dropped} direction sentences left out for mentioning logins; HASC link filled in for {hasc_filled} sites; {sum(len(v) for v in files_by_id.values())} attachments "
          f"(flow sends attachments: {'yes, ' + str(len(body.get('attachments') or [])) + ' items checked' if isinstance(body, dict) and 'attachments' in body else 'no'})")
    list_time_report(rows)
    return {"cols": cols, "rows": rows}


def main():
    secrets = json.loads(os.environ.get("ALL_SECRETS") or "{}")
    passcode = secrets.get("DASHBOARD_PASSCODE") or os.environ.get("DASHBOARD_PASSCODE")

    cols, rows = read_rows(get_file())
    ni = cols.index("Name")
    dropped = sum(1 for r in rows if SYSTEM_ACCOUNT.search(str(r[ni] or "")))
    rows = [r for r in rows if not SYSTEM_ACCOUNT.search(str(r[ni] or ""))]
    print(f"::notice title=Test accounts::Left out {dropped} HSI rows for system/test accounts (DoneSafe, Test Student).")
    if not rows:
        raise SystemExit("The Training Summary file had no data rows.")
    people = len({(r[0] or r[1]) for r in rows})
    print(f"{len(rows):,} training rows for {people:,} employees.")

    data = {"generated": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "cols": cols, "rows": rows}
    # Plant Sites tab: the TAPPISAFE attendance file (Tappi Raw Data), when its flow is set up.
    tappi_path = os.environ.get("TAPPI_FILE")
    tappi_bytes = open(tappi_path, "rb").read() if tappi_path else fetch_xlsx("TAPPI_URL", "TAPPI")
    tappi = read_tappi(tappi_bytes) if tappi_bytes else None
    if tappi:
        data["tappi"] = tappi
    # Refresher Training view: HSI Training Requirement Status (HSI Refresher Training.xlsx).
    ref_path = os.environ.get("REFRESHER_FILE")
    ref_bytes = open(ref_path, "rb").read() if ref_path else fetch_xlsx("REFRESHER_URL", "Refresher")
    refresher = read_refresher(ref_bytes) if ref_bytes else None
    if refresher:
        data["refresher"] = refresher
    if LENGTHS:
        data["courseLengths"] = LENGTHS
    site_reqs = get_site_requirements()
    if site_reqs:
        data["siteReqs"] = site_reqs
    shutil.rmtree(OUT, ignore_errors=True)
    os.makedirs(OUT)
    for name in ("index.html", "plants.js"):
        shutil.copy(os.path.join(HERE, name), os.path.join(OUT, name))
    if passcode:
        with open(os.path.join(OUT, "data.enc.json"), "w") as f:
            json.dump(encrypt(data, passcode), f)
        print("Data locked with DASHBOARD_PASSCODE.")
    else:
        with open(os.path.join(OUT, "data.json"), "w") as f:
            json.dump(data, f, separators=(",", ":"))
        print("No DASHBOARD_PASSCODE secret: publishing without a passcode.")
    with open(os.path.join(OUT, "robots.txt"), "w") as f:
        f.write("User-agent: *\nDisallow: /\n")
    print(f"::notice title=Training dashboard::{len(rows):,} rows, {people:,} employees"
          + (f", {len(tappi['rows']):,} TAPPI records" if tappi else ", no TAPPI file yet (plant-site readiness off; add the TAPPI_URL secret)")
          + (f", {len(refresher['rows']):,} refresher rows" if refresher else ", no refresher file yet (add the REFRESHER_URL secret)"))


if __name__ == "__main__":
    main()
